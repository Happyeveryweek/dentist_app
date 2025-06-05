import pandas as pd
import os
import json
from datetime import datetime
from PyQt6.QtWidgets import QFileDialog, QPushButton, QMessageBox, QTableWidget
from PyQt6.QtCore import Qt


def format_dental_condition(dental_condition_json):
    """
    将牙齿状况的JSON数据转换为易读的文本格式
    
    参数:
        dental_condition_json: 牙齿状况的JSON字符串
        
    返回:
        str: 格式化后的牙齿状况文本
    """
    if not dental_condition_json:
        return ""
        
    try:
        # 将JSON字符串转换为字典
        if isinstance(dental_condition_json, str):
            dental_data = json.loads(dental_condition_json)
        else:
            dental_data = dental_condition_json
            
        # 如果转换失败或为空，返回原始数据
        if not dental_data or not isinstance(dental_data, dict):
            return str(dental_condition_json)
            
        # 格式化输出
        result_lines = []
        
        # 获取所有日期键并排序
        date_keys = sorted([k for k in dental_data.keys() if k.startswith('date-')],
                          key=lambda x: int(x.split('-')[1]))
                          
        # 遍历每个日期键
        for date_key in date_keys:
            # 获取行索引
            row_idx = int(date_key.split('-')[1])
            # 获取日期值
            date_value = dental_data.get(date_key, "")
            
            # 添加日期行
            result_lines.append(f"日期-{row_idx}: {date_value}")
            
            # 遍历十字图表数据
            for chart_num in [1, 2]:
                result_lines.append(f"  图表{chart_num}:")
                
                # 定义位置映射
                position_map = {
                    "top-left": "左上",
                    "top-right": "右上",
                    "bottom-left": "左下",
                    "bottom-right": "右下"
                }
                
                # 遍历四个位置
                for pos, pos_cn in position_map.items():
                    # 构造键名
                    key = f"chart{chart_num}-{pos}-{row_idx}"
                    # 获取值
                    value = dental_data.get(key, "")
                    # 如果有值，添加到结果
                    if value:
                        result_lines.append(f"    {pos_cn}-{row_idx}: {value}")
        
        # 合并所有行
        return "\n".join(result_lines)
    except Exception as e:
        print(f"格式化牙齿状况时出错: {e}")
        return str(dental_condition_json)  # 失败时返回原始数据


def export_patients_to_excel(parent, patients_data, selected_only=False):
    """
    将患者信息导出到Excel文件
    
    参数:
        parent: 父窗口，用于显示文件对话框
        patients_data: 可以是患者表格控件(QTableWidget)或患者对象列表([Patient])
        selected_only: 是否只导出选中的患者
    """
    try:
        # 准备数据
        data = []
        # 更新表头，添加病历号、身份证号、医生字段
        headers = ["ID", "病历号", "姓名", "年龄", "性别", "电话", "地址", "身份证号", "医生", "初诊时间", "牙齿状况", "诊疗费用项目", "总费用", "创建时间", "更新时间"]
        
        # 判断输入类型，直接处理Patient对象列表
        if not isinstance(patients_data, QTableWidget):
            # 直接处理患者对象列表
            patients = patients_data
            for patient in patients:
                row_data = {}
                # 获取患者属性
                row_data["ID"] = str(patient.id)
                row_data["病历号"] = str(patient.medical_record_number) if patient.medical_record_number else ""
                row_data["姓名"] = patient.name
                row_data["年龄"] = str(patient.age) if patient.age else ""
                row_data["性别"] = patient.gender or ""
                row_data["电话"] = patient.phone or ""
                row_data["地址"] = patient.address or ""
                row_data["身份证号"] = patient.identification_number or ""
                row_data["医生"] = patient.doctor or ""
                row_data["初诊时间"] = patient.first_visit_date.strftime("%Y-%m-%d") if patient.first_visit_date else ""
                
                # 格式化牙齿状况
                row_data["牙齿状况"] = format_dental_condition(patient.dental_condition)
                
                row_data["诊疗费用项目"] = patient.treatment_items or ""
                row_data["总费用"] = f"¥{patient.total_cost:.2f}" if hasattr(patient, 'total_cost') and patient.total_cost else "¥0.00"
                row_data["创建时间"] = patient.created_at.strftime("%Y-%m-%d %H:%M:%S") if hasattr(patient, 'created_at') and patient.created_at else ""
                row_data["更新时间"] = patient.updated_at.strftime("%Y-%m-%d %H:%M:%S") if hasattr(patient, 'updated_at') and patient.updated_at else ""
                
                data.append(row_data)
        else:
            # 处理表格控件
            patients_table = patients_data
            # 获取选中的行
            selected_rows = []
            if selected_only:
                # 这里的selected_only参数已经在desktop_app_pyqt6.py中处理
                # 传入的patients_table已经是只包含选中行的临时表格
                selected_rows = list(range(patients_table.rowCount()))
            else:
                # 导出所有行
                selected_rows = list(range(patients_table.rowCount()))
            
            if not selected_rows:
                QMessageBox.warning(parent, "导出提示", "没有可导出的患者数据")
                return
            
            # 从表格中提取数据
            for row in selected_rows:
                row_data = {}
                # 从表格中获取可见数据
                # 注意：如果是选中导出，表格的列索引已经调整（没有选择列）
                if selected_only:
                    # 选中导出时，临时表格已经去掉了选择列，所以索引从0开始
                    id_col = 0  # ID列索引
                    name_col = 1  # 姓名列索引
                    age_col = 2  # 年龄列索引
                    gender_col = 3  # 性别列索引
                    phone_col = 4  # 电话列索引
                    first_visit_col = 5  # 初诊时间列索引
                    total_fee_col = 6  # 总费用列索引
                else:
                    # 正常导出时，如果选择列可见，则需要偏移1
                    is_selection_visible = not patients_table.isColumnHidden(0)
                    offset = 1 if is_selection_visible else 0
                    id_col = 0 + offset  # ID列索引
                    name_col = 1 + offset  # 姓名列索引
                    age_col = 2 + offset  # 年龄列索引
                    gender_col = 3 + offset  # 性别列索引
                    phone_col = 4 + offset  # 电话列索引
                    first_visit_col = 5 + offset  # 初诊时间列索引
                    total_fee_col = 6 + offset  # 总费用列索引
                
                # 获取表格数据
                row_data["ID"] = patients_table.item(row, id_col).text() if patients_table.item(row, id_col) else ""
                row_data["姓名"] = patients_table.item(row, name_col).text() if patients_table.item(row, name_col) else ""
                row_data["年龄"] = patients_table.item(row, age_col).text() if patients_table.item(row, age_col) else ""
                row_data["性别"] = patients_table.item(row, gender_col).text() if patients_table.item(row, gender_col) else ""
                row_data["电话"] = patients_table.item(row, phone_col).text() if patients_table.item(row, phone_col) else ""
                row_data["初诊时间"] = patients_table.item(row, first_visit_col).text() if patients_table.item(row, first_visit_col) else ""
                row_data["总费用"] = patients_table.item(row, total_fee_col).text() if patients_table.item(row, total_fee_col) else ""
                
                # 其他字段设为空，后续可以从数据库中获取完整信息
                row_data["地址"] = ""
                row_data["牙齿状况"] = ""  # 从表格中无法获取牙齿状况的详细数据
                row_data["诊疗费用项目"] = ""
                row_data["创建时间"] = ""
                row_data["更新时间"] = ""
                
                data.append(row_data)
        
        # 如果没有数据
        if not data:
            QMessageBox.warning(parent, "导出提示", "没有可导出的患者数据")
            return
            
        # 创建DataFrame
        df = pd.DataFrame(data)
        
        # 获取保存路径
        current_time = datetime.now().strftime("%Y%m%d_%H%M%S")
        default_filename = f"患者信息_{current_time}.xlsx"
        file_path, _ = QFileDialog.getSaveFileName(
            parent,
            "保存患者信息",
            os.path.join(os.path.expanduser("~"), "Documents", default_filename),
            "Excel文件 (*.xlsx);;所有文件 (*)"
        )
        
        if not file_path:
            return  # 用户取消了保存
        
        # 如果用户没有指定.xlsx扩展名，添加它
        if not file_path.endswith(".xlsx"):
            file_path += ".xlsx"
        
        # 保存到Excel
        df.to_excel(file_path, index=False, sheet_name="患者信息")
        
        # 使用openpyxl调整Excel格式
        from openpyxl import load_workbook
        from openpyxl.styles import Font, Alignment, PatternFill
        
        # 加载工作簿
        wb = load_workbook(file_path)
        ws = wb.active
        
        # 设置列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter  # 获取列字母
            column_index = col[0].column  # 获取列索引
            
            # 牙齿状况列特殊处理
            if headers[column_index-1] == "牙齿状况":  # Excel列索引从1开始，而headers是从0开始
                # 为牙齿状况列设置固定宽度，避免过宽
                adjusted_width = 30  # 固定设置为30个字符宽度
            # 身份证号列特殊处理
            elif headers[column_index-1] == "身份证号":
                # 设置合适的宽度
                for cell in col:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                # 调整宽度
                adjusted_width = (max_length + 2) * 1.2
                # 对该列中的每个单元格应用文本格式和居中对齐
                for cell in col:
                    if cell.row > 1:  # 跳过表头
                        # 设置为文本格式
                        cell.number_format = '@'
                        # 设置上下左右居中对齐
                        cell.alignment = Alignment(horizontal="center", vertical="center")
            # 电话号码列特殊处理
            elif headers[column_index-1] == "电话":
                # 设置合适的宽度
                for cell in col:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                # 调整宽度
                adjusted_width = (max_length + 2) * 1.2
                # 对该列中的每个单元格应用文本格式和居中对齐
                for cell in col:
                    if cell.row > 1:  # 跳过表头
                        # 设置为文本格式
                        cell.number_format = '@'
                        # 设置上下左右居中对齐
                        cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                # 其他列正常处理
                for cell in col:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                # 标准宽度调整
                adjusted_width = (max_length + 2) * 1.2
            
            # 设置列宽
            ws.column_dimensions[column].width = adjusted_width
        
        # 设置表头样式
        header_fill = PatternFill(start_color="4A86E8", end_color="4A86E8", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # 设置数据行样式
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                column_title = headers[cell.column-1]
                if column_title == "牙齿状况":  # 牙齿状况列
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                elif column_title == "身份证号":  # 身份证号列
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.number_format = '@'  # 设置为文本格式
                elif column_title == "电话":  # 电话号码列
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.number_format = '@'  # 设置为文本格式
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # 保存修改后的工作簿
        wb.save(file_path)
        
        QMessageBox.information(parent, "导出成功", f"患者信息已成功导出到:\n{file_path}")
        
    except Exception as e:
        QMessageBox.critical(parent, "导出错误", f"导出患者信息时出错:\n{str(e)}")
        print(f"导出患者信息时出错: {e}")


def create_export_button():
    """
    创建导出按钮
    """
    export_btn = QPushButton("导出Excel")
    export_btn.setCursor(Qt.CursorShape.PointingHandCursor)
    export_btn.setStyleSheet("""
        QPushButton {
            background-color: #4CAF50;
            color: white;
            border: none;
            border-radius: 8px;
            padding: 12px 24px;
            font-weight: 600;
            font-size: 14px;
        }
        QPushButton:hover {
            background-color: #45a049;
        }
        QPushButton:pressed {
            background-color: #3d8b40;
        }
    """)
    return export_btn